"""Cross-check the MWA membership spreadsheet against the MWA ASVO user list.

The script reads one worksheet of the membership workbook and gets the first name, surname, institution and email
address of each member. It also gets alternate email addresses from the "Notes" column. Then it runs an SQL query
on the MWA ASVO database and matches each member to an ASVO user. It matches by email first, then by name.

The SQL query must return these columns (use ``AS`` aliases if necessary): ``first_name``, ``last_name``, ``email``.

The database connection uses a libpq connection string. If you do not supply one, libpq uses the standard ``PG*``
environment variables and ``~/.pgpass``.

Example:
    python mwa_member_check.py members.xlsx --sheet "Members" --sql-file asvo_users.sql
"""

import argparse
import re
import sys
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import LiteralString, cast

import openpyxl
import psycopg
from psycopg.rows import dict_row

# Spreadsheet column headers (compared after whitespace is stripped).
COL_FIRST_NAME = "First Name"
COL_LAST_NAME = "Last Name"
COL_INSTITUTION = "Institution"
COL_EMAIL = "Email"
COL_NOTES = "Notes"
REQUIRED_COLUMNS = (COL_FIRST_NAME, COL_LAST_NAME, COL_INSTITUTION, COL_EMAIL, COL_NOTES)

# The header is on this row of the worksheet (1-based).
HEADER_ROW = 1

# Column names that the ASVO SQL query must return.
SQL_COL_FIRST_NAME = "first_name"
SQL_COL_LAST_NAME = "last_name"
SQL_COL_EMAIL = "email"
REQUIRED_SQL_COLUMNS = (SQL_COL_FIRST_NAME, SQL_COL_LAST_NAME, SQL_COL_EMAIL)

# Finds email addresses in free text. Commas, spaces and semicolons are not part of a match, so they act as separators.
EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

# Unicode normalisation form for name comparison.
UNICODE_FORM = "NFKC"

# Default libpq connection string. Empty means: use the PG* environment variables and ~/.pgpass.
DEFAULT_DSN = ""

# Report layout.
SECTION_RULE = "=" * 80
FIELD_SEPARATOR = " | "


@dataclass(frozen=True)
class Member:
    """One MWA member from the membership spreadsheet.

    Attributes:
        row: Worksheet row number (1-based).
        first_name: First name.
        last_name: Surname.
        institution: Institution.
        email: Primary email address.
        alt_emails: Alternate email addresses from the Notes column.
    """

    row: int
    first_name: str
    last_name: str
    institution: str
    email: str
    alt_emails: tuple[str, ...] = field(default_factory=tuple)

    def all_emails(self) -> list[str]:
        """Return the primary and alternate email addresses, normalised, without duplicates.

        Returns:
            Normalised email addresses. The primary address is first.
        """
        result: list[str] = []
        for address in (self.email, *self.alt_emails):
            key = normalise_email(address)
            if key and key not in result:
                result.append(key)
        return result

    def describe(self) -> str:
        """Return a one-line description for the report.

        Returns:
            The member details on one line.
        """
        parts = [f"row {self.row}", f"{self.first_name} {self.last_name}", self.institution, self.email or "(no email)"]
        if self.alt_emails:
            parts.append("alt: " + ", ".join(self.alt_emails))
        return FIELD_SEPARATOR.join(parts)


@dataclass(frozen=True)
class AsvoUser:
    """One user from the MWA ASVO database.

    Attributes:
        first_name: First name.
        last_name: Surname.
        email: Email address.
    """

    first_name: str
    last_name: str
    email: str

    def describe(self) -> str:
        """Return a one-line description for the report.

        Returns:
            The user details on one line.
        """
        return FIELD_SEPARATOR.join([f"{self.first_name} {self.last_name}", self.email or "(no email)"])


@dataclass
class MatchResult:
    """The result of the cross-check.

    Attributes:
        by_email: Members that match an ASVO user by email.
        by_name: Members that match one or more ASVO users by name only.
        unmatched_members: Members with no match.
        unmatched_users: ASVO users that match no member.
    """

    by_email: list[tuple[Member, AsvoUser]] = field(default_factory=list)
    by_name: list[tuple[Member, list[AsvoUser]]] = field(default_factory=list)
    unmatched_members: list[Member] = field(default_factory=list)
    unmatched_users: list[AsvoUser] = field(default_factory=list)


def clean_text(value: object) -> str:
    """Convert a cell or column value to a stripped string.

    Args:
        value: The raw value. Can be None.

    Returns:
        The stripped string. An empty string if the value is None.
    """
    if value is None:
        return ""
    return str(value).strip()


def normalise_email(address: str) -> str:
    """Normalise an email address for comparison.

    Args:
        address: The email address.

    Returns:
        The stripped, lower-case address.
    """
    return address.strip().casefold()


def normalise_name(first_name: str, last_name: str) -> tuple[str, str]:
    """Normalise a name for comparison.

    Applies Unicode normalisation, case folding and whitespace collapse.

    Args:
        first_name: First name.
        last_name: Surname.

    Returns:
        The normalised (first name, surname) pair.
    """

    def _norm(text: str) -> str:
        return " ".join(unicodedata.normalize(UNICODE_FORM, text).casefold().split())

    return _norm(first_name), _norm(last_name)


def extract_emails(text: str) -> list[str]:
    """Find all email addresses in free text.

    Args:
        text: The text. Addresses can be separated by commas, spaces or semicolons.

    Returns:
        The email addresses in the order they occur.
    """
    return EMAIL_PATTERN.findall(text)


def read_members(path: Path, sheet_name: str) -> list[Member]:
    """Read the members from one worksheet of the membership workbook.

    Rows with no first name, no surname and no email are ignored.

    Args:
        path: Path to the .xlsx file.
        sheet_name: Name of the worksheet.

    Returns:
        The members in worksheet order.

    Raises:
        KeyError: If the worksheet does not exist.
        ValueError: If a required column header is missing.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"Worksheet {sheet_name!r} not found. Worksheets: {workbook.sheetnames}")
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(min_row=HEADER_ROW, values_only=True)
        header = [clean_text(cell) for cell in next(rows, ())]
        missing = [name for name in REQUIRED_COLUMNS if name not in header]
        if missing:
            raise ValueError(f"Missing column(s) {missing} in header row {HEADER_ROW}. Found: {header}")
        index = {name: header.index(name) for name in REQUIRED_COLUMNS}

        members: list[Member] = []
        for row_number, values in enumerate(rows, start=HEADER_ROW + 1):

            def _get(column: str, cells: tuple[object, ...] = values) -> str:
                position = index[column]
                return clean_text(cells[position]) if position < len(cells) else ""

            first_name = _get(COL_FIRST_NAME)
            last_name = _get(COL_LAST_NAME)
            email = _get(COL_EMAIL)
            if not (first_name or last_name or email):
                continue
            primary = normalise_email(email)
            alt_emails = tuple(
                address for address in extract_emails(_get(COL_NOTES)) if normalise_email(address) != primary
            )
            members.append(
                Member(
                    row=row_number,
                    first_name=first_name,
                    last_name=last_name,
                    institution=_get(COL_INSTITUTION),
                    email=email,
                    alt_emails=alt_emails,
                )
            )
        return members
    finally:
        workbook.close()


def read_asvo_users(dsn: str, query: str) -> list[AsvoUser]:
    """Run the ASVO user query and return the users.

    Args:
        dsn: libpq connection string. Empty means: use the PG* environment variables and ~/.pgpass.
        query: SQL query. It must return the columns in REQUIRED_SQL_COLUMNS.

    Returns:
        The ASVO users.

    Raises:
        ValueError: If the query does not return a required column.
    """
    # The query comes from a local file that the operator supplies, so it is trusted. psycopg types queries as
    # LiteralString to prevent SQL injection; the cast tells the type checker that this query is trusted.
    trusted_query = cast("LiteralString", query)
    with psycopg.connect(dsn) as conn, conn.cursor(row_factory=dict_row) as cursor:
        cursor.execute(trusted_query)
        columns = [column.name for column in cursor.description or []]
        missing = [name for name in REQUIRED_SQL_COLUMNS if name not in columns]
        if missing:
            raise ValueError(f"SQL query does not return column(s) {missing}. Returned: {columns}")
        return [
            AsvoUser(
                first_name=clean_text(row[SQL_COL_FIRST_NAME]),
                last_name=clean_text(row[SQL_COL_LAST_NAME]),
                email=clean_text(row[SQL_COL_EMAIL]),
            )
            for row in cursor.fetchall()
        ]


def cross_check(members: Iterable[Member], users: Iterable[AsvoUser]) -> MatchResult:
    """Match members to ASVO users. Email match first, name match as fallback.

    Args:
        members: The members from the spreadsheet.
        users: The ASVO users.

    Returns:
        The match result.
    """
    user_list = list(users)
    by_email: dict[str, AsvoUser] = {}
    by_name: dict[tuple[str, str], list[AsvoUser]] = {}
    for user in user_list:
        email_key = normalise_email(user.email)
        if email_key:
            by_email.setdefault(email_key, user)
        by_name.setdefault(normalise_name(user.first_name, user.last_name), []).append(user)

    result = MatchResult()
    matched_ids: set[int] = set()
    for member in members:
        email_match = next((by_email[key] for key in member.all_emails() if key in by_email), None)
        if email_match is not None:
            result.by_email.append((member, email_match))
            matched_ids.add(id(email_match))
            continue
        name_matches = by_name.get(normalise_name(member.first_name, member.last_name), [])
        if name_matches:
            result.by_name.append((member, name_matches))
            matched_ids.update(id(user) for user in name_matches)
            continue
        result.unmatched_members.append(member)

    result.unmatched_users = [user for user in user_list if id(user) not in matched_ids]
    return result


def print_report(result: MatchResult, show_unmatched_users: bool) -> None:
    """Print the cross-check report to stdout.

    Args:
        result: The match result.
        show_unmatched_users: If True, list each ASVO user that matches no member. If False, print only the count.
    """
    print(SECTION_RULE)
    print(f"Matched by email: {len(result.by_email)}")
    print(SECTION_RULE)
    for member, user in result.by_email:
        print(f"  {member.describe()}")
        print(f"    -> ASVO: {user.describe()}")

    print(SECTION_RULE)
    print(f"Matched by name only (check the email): {len(result.by_name)}")
    print(SECTION_RULE)
    for member, users in result.by_name:
        print(f"  {member.describe()}")
        for user in users:
            print(f"    -> ASVO: {user.describe()}")

    print(SECTION_RULE)
    print(f"Members with no ASVO account: {len(result.unmatched_members)}")
    print(SECTION_RULE)
    for member in result.unmatched_members:
        print(f"  {member.describe()}")

    print(SECTION_RULE)
    print(f"ASVO users that match no member: {len(result.unmatched_users)}")
    print(SECTION_RULE)
    if show_unmatched_users:
        for user in result.unmatched_users:
            print(f"  {user.describe()}")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse the command-line arguments.

    Args:
        argv: The arguments. None means sys.argv.

    Returns:
        The parsed arguments.
    """
    parser = argparse.ArgumentParser(description="Cross-check the MWA membership spreadsheet against MWA ASVO users.")
    parser.add_argument("workbook", type=Path, help="Path to the membership .xlsx file.")
    parser.add_argument("--sheet", required=True, help="Name of the worksheet to read.")
    parser.add_argument(
        "--sql-file",
        type=Path,
        required=True,
        help=f"File with the ASVO user SQL query. It must return: {', '.join(REQUIRED_SQL_COLUMNS)}.",
    )
    parser.add_argument(
        "--dsn",
        default=DEFAULT_DSN,
        help="libpq connection string. Default: use the PG* environment variables and ~/.pgpass.",
    )
    parser.add_argument(
        "--show-unmatched-users",
        action="store_true",
        help="List each ASVO user that matches no member. Default: print only the count.",
    )
    parser.add_argument(
        "--members-only",
        action="store_true",
        help="Print only the member list from the spreadsheet. Do not connect to the database.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """Run the cross-check.

    Args:
        argv: The arguments. None means sys.argv.

    Returns:
        The process exit code.
    """
    args = parse_args(argv)
    members = read_members(args.workbook, args.sheet)
    if args.members_only:
        for member in members:
            print(member.describe())
        print(f"Total members: {len(members)}")
        return 0

    query = args.sql_file.read_text(encoding="utf-8")
    users = read_asvo_users(args.dsn, query)
    print(f"Members: {len(members)}  ASVO users: {len(users)}")
    print_report(cross_check(members, users), args.show_unmatched_users)
    return 0


if __name__ == "__main__":
    sys.exit(main())
